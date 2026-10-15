"""Carregamento de planilhas CSV/XLSX com schema compartilhado e cache.

Fluxo para XLSX: ``get_schema`` (xldetect, em cache) -> ``read_workbook``
(openpyxl, guiado pelo schema) -> registros em cache. Os registros em
cache são compartilhados: quem os recebe não deve alterá-los.
"""

from __future__ import annotations

import csv
import threading
import warnings
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .schema import RegionSchema, WorkbookSchema, build_schema

CACHE_MAX_ENTRIES = 16

# (caminho resolvido, mtime em ns, tamanho em bytes)
CacheKey = tuple[str, int, int]

Records = list[dict] | dict[str, list[dict]]


class _BoundedCache:
    """Cache em memória limitado, com descarte do mais antigo (FIFO)."""

    def __init__(self, max_entries: int) -> None:
        self._max_entries = max_entries
        self._items: dict[CacheKey, Any] = {}
        self._lock = threading.RLock()

    def get(self, key: CacheKey) -> Any | None:
        with self._lock:
            return self._items.get(key)

    def put(self, key: CacheKey, value: Any) -> None:
        with self._lock:
            # Entradas antigas do mesmo arquivo ficam obsoletas.
            for stale in [
                k for k in self._items if k[0] == key[0] and k != key
            ]:
                del self._items[stale]

            while len(self._items) >= self._max_entries:
                self._items.pop(next(iter(self._items)))

            self._items[key] = value

    def clear(self) -> None:
        with self._lock:
            self._items.clear()

    def __len__(self) -> int:
        with self._lock:
            return len(self._items)


_SCHEMAS = _BoundedCache(CACHE_MAX_ENTRIES)
_RECORDS = _BoundedCache(CACHE_MAX_ENTRIES)

# Evita que leituras concorrentes do mesmo arquivo (ainda não cacheado)
# sejam feitas em duplicidade. A granularidade é por arquivo.
_CACHE_LOCKS: dict[str, threading.Lock] = {}
_CACHE_LOCKS_GUARD = threading.Lock()


def _lock_for(path: Path) -> threading.Lock:
    chave = str(path)

    with _CACHE_LOCKS_GUARD:
        lock = _CACHE_LOCKS.get(chave)

        if lock is None:
            lock = threading.RLock()
            _CACHE_LOCKS[chave] = lock

        return lock


def _file_key(path: Path) -> CacheKey:
    stat = path.stat()
    return (str(path), stat.st_mtime_ns, stat.st_size)


def clear_cache() -> None:
    """Esvazia os caches de schema e de registros."""
    _SCHEMAS.clear()
    _RECORDS.clear()


def read_csv(path: str | Path) -> list[dict]:
    path = Path(path)
    encodings = ["utf-8-sig", "utf-8", "cp1252", "latin-1"]

    for encoding in encodings:
        try:
            with path.open(
                "r",
                encoding=encoding,
                newline="",
            ) as arquivo:
                return list(csv.DictReader(arquivo))
        except UnicodeDecodeError:
            continue

    raise ValueError("Não foi possível ler o arquivo")


def _read_region(ws, region: RegionSchema) -> list[dict]:
    """Extrai os registros de uma região descrita pelo schema."""
    max_row = region["max_row"]
    min_col = region["min_col"]
    max_col = region["max_col"]
    first_data_row = region["first_data_row"]
    headers = region["columns"]

    # Reconstrói apenas células pertencentes a mesclagens verticais
    # dentro desta região. Não preenche vazios comuns.
    merged_values = {}

    for merged in ws.merged_cells.ranges:
        if merged.min_col != merged.max_col:
            continue

        if not (
            min_col <= merged.min_col <= max_col
            and first_data_row <= merged.min_row <= max_row
            and merged.max_row <= max_row
        ):
            continue

        value = ws.cell(
            merged.min_row,
            merged.min_col,
        ).value

        for row_num in range(
            merged.min_row,
            merged.max_row + 1,
        ):
            merged_values[(row_num, merged.min_col)] = value

    records = []

    for row_num in range(first_data_row, max_row + 1):
        values = []

        for col_num in range(min_col, max_col + 1):
            value = merged_values.get(
                (row_num, col_num),
                ws.cell(row_num, col_num).value,
            )
            values.append(value)

        if not any(value is not None for value in values):
            continue

        records.append(dict(zip(headers, values)))

    return records


def read_workbook(
    path: str | Path,
    schema: WorkbookSchema,
) -> dict[str, list[dict]]:
    """Lê os registros de um XLSX usando um schema já construído."""
    wb = load_workbook(
        path,
        read_only=False,
        data_only=True,
    )

    try:
        data = {}

        for ws in wb.worksheets:
            regions = schema["sheets"].get(ws.title, [])
            records = []

            for region in regions:
                records.extend(_read_region(ws, region))

            if not regions:
                # Preserva conteúdo caso o detector não encontre regiões.
                warnings.warn(
                    f"xldetect não encontrou regiões na aba "
                    f"{ws.title!r}; preservando células não vazias "
                    f"com nomes genéricos."
                )

                for row in ws.iter_rows():
                    values = [cell.value for cell in row]

                    if not any(value is not None for value in values):
                        continue

                    records.append({
                        f"coluna_{index}": value
                        for index, value in enumerate(values, start=1)
                    })

            data[ws.title] = records

        return data
    finally:
        wb.close()


def get_schema(path: str | Path) -> WorkbookSchema:
    """Schema normalizado de um XLSX (em cache até o arquivo mudar)."""
    path = Path(path).resolve()

    if path.suffix.lower() != ".xlsx":
        raise ValueError("Schema só está disponível para arquivos XLSX")

    key = _file_key(path)
    schema = _SCHEMAS.get(key)

    if schema is None:
        with _lock_for(path):
            key = _file_key(path)
            schema = _SCHEMAS.get(key)

            if schema is None:
                schema = build_schema(path)
                _SCHEMAS.put(key, schema)

    return schema


def load_records(path: str | Path) -> Records:
    """Registros de um CSV (lista) ou XLSX (dict aba -> lista).

    O resultado fica em cache até o arquivo mudar; não o modifique.
    """
    path = Path(path).resolve()
    suffix = path.suffix.lower()

    if suffix not in (".csv", ".xlsx"):
        raise ValueError("Extensão não suportada")

    key = _file_key(path)
    records = _RECORDS.get(key)

    if records is None:
        with _lock_for(path):
            key = _file_key(path)
            records = _RECORDS.get(key)

            if records is None:
                if suffix == ".csv":
                    records = read_csv(path)
                else:
                    records = read_workbook(path, get_schema(path))

                _RECORDS.put(key, records)

    return records
