"""Descoberta da estrutura de planilhas XLSX (schema normalizado).

O schema descreve *onde* estão as tabelas de cada aba e *quais* são suas
colunas. Ele não contém registros. Estrutura:

    {
        "path": "/caminho/absoluto/arquivo.xlsx",
        "sheets": {                      # ordem original das abas
            "Aba": [                     # ordem original das regiões
                {
                    "sheet": "Aba",
                    "min_row": 1, "max_row": 10,
                    "min_col": 1, "max_col": 4,
                    "has_header": True,
                    "header_row": 1,       # None quando não há cabeçalho
                    "first_data_row": 2,
                    "columns": ["Nome", "coluna_2", "Nome_2", "Valor"],
                },
            ],
            "Aba sem regiões": [],       # o xldetect não encontrou nada
        },
    }

Linhas e colunas são 1-based e inclusivas, como no Excel.
"""

from __future__ import annotations

from collections.abc import Iterable
from pathlib import Path
from typing import TypedDict

from openpyxl import load_workbook
from xldetect import inspect_path


class RegionSchema(TypedDict):
    sheet: str
    min_row: int
    max_row: int
    min_col: int
    max_col: int
    has_header: bool
    header_row: int | None
    first_data_row: int
    columns: list[str]


class WorkbookSchema(TypedDict):
    path: str
    sheets: dict[str, list[RegionSchema]]


def make_headers(values: Iterable[object], start_col: int) -> list[str]:
    """Gera cabeçalhos válidos, preenchendo vazios e evitando duplicatas."""
    headers: list[str] = []
    used: set[str] = set()

    for offset, value in enumerate(values):
        name = str(value).strip() if value is not None else ""

        if not name:
            name = f"coluna_{start_col + offset}"

        base = name
        suffix = 2

        while name in used:
            name = f"{base}_{suffix}"
            suffix += 1

        headers.append(name)
        used.add(name)

    return headers


def _read_header_cells(
    path: Path,
    sheet: str,
    row: int,
    min_col: int,
    max_col: int,
) -> list[object]:
    """Lê somente as células da linha de cabeçalho (sem ler registros)."""
    wb = load_workbook(path, read_only=True, data_only=True)

    try:
        rows = wb[sheet].iter_rows(
            min_row=row,
            max_row=row,
            min_col=min_col,
            max_col=max_col,
            values_only=True,
        )
        values = list(next(rows, ()))
    finally:
        wb.close()

    width = max_col - min_col + 1
    return values + [None] * (width - len(values))


def _normalize_region(path: Path, sheet: str, region) -> RegionSchema:
    """Converte uma região do xldetect em um dicionário simples.

    Atributos opcionais do xldetect são lidos com ``getattr`` e valores
    ausentes ou vazios caem nos mesmos padrões que o extrator usava.
    """
    min_row = int(region.min_row)
    max_row = int(region.max_row)
    min_col = int(region.min_col)
    max_col = int(region.max_col)
    width = max_col - min_col + 1

    header_row = getattr(region, "header_row", None)
    has_header = bool(getattr(region, "has_header", False))

    # Região com has_header mas sem linha de cabeçalho identificada é
    # tratada como região sem cabeçalho, igual ao extrator anterior.
    if not has_header or header_row is None:
        has_header = False
        header_row = None

    if header_row is not None:
        header_row = int(header_row)
        data_start = int(getattr(region, "data_start_row", 0) or 0)
        first_data_row = data_start or header_row + 1

        detected = list(getattr(region, "headers", None) or [])

        if len(detected) == width:
            values: list[object] = detected
        else:
            values = _read_header_cells(
                path, sheet, header_row, min_col, max_col
            )
    else:
        # Sem cabeçalho: todas as linhas da região são dados.
        first_data_row = min_row
        values = [None] * width

    return {
        "sheet": sheet,
        "min_row": min_row,
        "max_row": max_row,
        "min_col": min_col,
        "max_col": max_col,
        "has_header": has_header,
        "header_row": header_row,
        "first_data_row": first_data_row,
        "columns": make_headers(values, min_col),
    }


def build_schema(path: str | Path) -> WorkbookSchema:
    """Detecta as regiões tabulares de um XLSX e devolve o schema."""
    path = Path(path).resolve()
    report = inspect_path(str(path))

    sheets: dict[str, list[RegionSchema]] = {}

    for sheet_report in report.sheets:
        sheets[sheet_report.sheet] = [
            _normalize_region(path, sheet_report.sheet, region)
            for region in sheet_report.regions
        ]

    return {"path": str(path), "sheets": sheets}
