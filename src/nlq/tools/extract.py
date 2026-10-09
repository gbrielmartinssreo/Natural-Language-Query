from pathlib import Path
import csv
import json
import time
import warnings

from openpyxl import load_workbook
from xldetect import inspect_path

from langchain_core.tools import tool


BASE_DIR = Path(__file__).resolve().parent

_CACHE: dict[tuple[str, float, int], str] = {}
_CACHE_MAX_ENTRIES = 8


def _perf(label: str, seconds: float | None = None) -> None:
    if seconds is None:
        print(f"[PERF] {label}", flush=True)
    else:
        print(f"[PERF] {label}: {seconds:.2f}s", flush=True)


def _cache_key(path: Path) -> tuple[str, float, int]:
    stat = path.stat()
    return (str(path.resolve()), stat.st_mtime, stat.st_size)


def _store_cache(key: tuple[str, float, int], value: str) -> None:
    path_key = key[0]

    for stale in [
        k for k in _CACHE
        if k[0] == path_key and k != key
    ]:
        del _CACHE[stale]

    while len(_CACHE) >= _CACHE_MAX_ENTRIES:
        _CACHE.pop(next(iter(_CACHE)))

    _CACHE[key] = value


def read_csv(path: Path) -> list[dict]:
    encodings = ["utf-8-sig", "utf-8", "cp1252", "latin-1"]

    for encoding in encodings:
        try:
            with path.open(
                "r",
                encoding=encoding,
                newline="",
            ) as arquivo:
                return list(csv.DictReader(arquivo))
        except UnicodeDecodeError:
            continue

    raise ValueError("Não foi possível ler o arquivo")


def _make_headers(values, start_col: int) -> list[str]:
    """Gera cabeçalhos válidos, preenchendo vazios e evitando duplicatas."""
    headers = []
    used = set()

    for offset, value in enumerate(values):
        name = str(value).strip() if value is not None else ""

        if not name:
            name = f"coluna_{start_col + offset}"

        base = name
        suffix = 2

        while name in used:
            name = f"{base}_{suffix}"
            suffix += 1

        headers.append(name)
        used.add(name)

    return headers


def _read_region(ws, region) -> list[dict]:
    """Extrai uma região detectada pelo xldetect."""
    min_row = region.min_row
    max_row = region.max_row
    min_col = region.min_col
    max_col = region.max_col

    has_header = bool(region.has_header)
    header_row = region.header_row if has_header else None

    if has_header and header_row is not None:
        first_data_row = region.data_start_row or header_row + 1

        detected_headers = list(region.headers or [])
        expected_width = max_col - min_col + 1

        if len(detected_headers) == expected_width:
            headers = _make_headers(detected_headers, min_col)
        else:
            values = [
                ws.cell(header_row, col).value
                for col in range(min_col, max_col + 1)
            ]
            headers = _make_headers(values, min_col)
    else:
        # Sem cabeçalho detectado, não inventamos que a primeira
        # linha seja um cabeçalho: preservamos todas como dados.
        first_data_row = min_row
        headers = [
            f"coluna_{col}"
            for col in range(min_col, max_col + 1)
        ]

    # Reconstrói apenas células pertencentes a mesclagens verticais
    # dentro desta região. Não faz fill-down de vazios comuns.
    merged_values = {}

    for merged in ws.merged_cells.ranges:
        if merged.min_col != merged.max_col:
            continue

        if not (
            min_col <= merged.min_col <= max_col
            and first_data_row <= merged.min_row <= max_row
            and merged.max_row <= max_row
        ):
            continue

        value = ws.cell(
            merged.min_row,
            merged.min_col,
        ).value

        for row_num in range(
            merged.min_row,
            merged.max_row + 1,
        ):
            merged_values[(row_num, merged.min_col)] = value

    records = []

    for row_num in range(first_data_row, max_row + 1):
        values = []

        for col_num in range(min_col, max_col + 1):
            value = merged_values.get(
                (row_num, col_num),
                ws.cell(row_num, col_num).value,
            )
            values.append(value)

        if not any(value is not None for value in values):
            continue

        records.append(dict(zip(headers, values)))

    return records


def read_excel(path: Path) -> dict[str, list[dict]]:
    # A detecção é executada sobre o arquivo, não por uma heurística
    # própria de cabeçalho.
    t_detect = time.perf_counter()
    report = inspect_path(str(path))
    _perf("xldetect inspect", time.perf_counter() - t_detect)

    regions_by_sheet = {}

    for region in report.iter_regions():
        regions_by_sheet.setdefault(region.sheet, []).append(region)

    t_open = time.perf_counter()
    wb = load_workbook(
        path,
        read_only=False,
        data_only=True,
    )
    _perf("open workbook", time.perf_counter() - t_open)

    try:
        t_parse = time.perf_counter()
        data = {}

        for ws in wb.worksheets:
            regions = regions_by_sheet.get(ws.title, [])
            records = []

            for region in regions:
                region_records = _read_region(ws, region)
                records.extend(region_records)

                _perf(
                    f"sheet {ws.title}, região {region.range_a1}: "
                    f"{len(region_records)} registros "
                    f"(confiança {region.confidence:.2f})"
                )

            if not regions:
                # Fallback conservador: não descarta conteúdo só porque
                # o detector não reconheceu uma região tabular.
                warnings.warn(
                    f"xldetect não encontrou regiões na aba "
                    f"{ws.title!r}; preservando células não vazias "
                    f"com nomes genéricos."
                )

                for row in ws.iter_rows():
                    values = [cell.value for cell in row]

                    if not any(value is not None for value in values):
                        continue

                    records.append({
                        f"coluna_{index}": value
                        for index, value in enumerate(values, start=1)
                    })

            data[ws.title] = records

        _perf("parse workbook", time.perf_counter() - t_parse)
        return data
    finally:
        wb.close()


@tool
def create_json(name: str, ext: str) -> str:
    """
    Lê uma planilha suportada e retorna seu conteúdo em JSON.

    Args:
        name: Nome do arquivo sem extensão, por exemplo academia.
        ext: Extensão do arquivo, por exemplo csv ou xlsx.
    """
    t_total = time.perf_counter()

    ext = ext.lower().lstrip(".")

    if ext not in ("csv", "xlsx"):
        raise ValueError("Extensão não suportada")

    path = (BASE_DIR / "../../../sheets" / ext / f"{name}.{ext}").resolve()

    if not path.is_file():
        raise FileNotFoundError(f"Arquivo não encontrado: {path}")

    key = _cache_key(path)
    cached = _CACHE.get(key)

    if cached is not None:
        _perf("cache hit")
        _perf("create_json total", time.perf_counter() - t_total)
        return cached

    if ext == "csv":
        t_read = time.perf_counter()
        dados = read_csv(path)
        _perf("read csv", time.perf_counter() - t_read)
    else:
        dados = read_excel(path)

    t_json = time.perf_counter()
    json_dados = json.dumps(
        dados,
        ensure_ascii=False,
        indent=2,
        default=str,
    )
    _perf("to json", time.perf_counter() - t_json)

    _store_cache(key, json_dados)
    _perf("create_json total", time.perf_counter() - t_total)

    return json_dados


if __name__ == "__main__":
    print(create_json("matriz", "xlsx"))
