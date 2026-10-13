from pathlib import Path
import json

from langchain_core.tools import tool
from openpyxl.utils import get_column_letter

from nlq.data import get_schema
from nlq.paths import data_dir


def _resolve_path(name: str, ext: str) -> Path:
    sheets_dir = data_dir().resolve()
    path = (sheets_dir / ext / f"{name}.{ext}").resolve()

    # O nome vem da LLM: não permite sair da pasta de planilhas.
    if not path.is_relative_to(sheets_dir):
        raise ValueError("Nome de arquivo inválido")

    if not path.is_file():
        raise FileNotFoundError(f"Arquivo não encontrado: {path}")

    return path


def _select_sheets(schema: dict, aba: str | None) -> dict[str, list[dict]]:
    sheets = schema["sheets"]

    if aba is None or not aba.strip():
        return sheets

    if aba in sheets:
        return {aba: sheets[aba]}

    matches = [
        title for title in sheets
        if title.strip().casefold() == aba.strip().casefold()
    ]

    if len(matches) == 1:
        return {matches[0]: sheets[matches[0]]}

    disponiveis = ", ".join(repr(title) for title in sheets)
    raise ValueError(
        f"Aba {aba!r} não encontrada. Abas disponíveis: {disponiveis}"
    )


def _plural(n: int, singular: str, plural: str) -> str:
    return f"{n} {singular if n == 1 else plural}"


def _format_region(index: int, region: dict) -> tuple[list[str], bool]:
    first = f"{get_column_letter(region['min_col'])}{region['min_row']}"
    last = f"{get_column_letter(region['max_col'])}{region['max_row']}"
    columns = region["columns"]

    if region["has_header"]:
        header = f"cabeçalho na linha {region['header_row']}"
    else:
        header = "sem cabeçalho detectado"

    lines = [
        f"  Região {index}: {first}:{last} "
        f"(linhas {region['min_row']} a {region['max_row']}, "
        f"colunas {region['min_col']} a {region['max_col']})",
        f"    {header}; dados a partir da linha {region['first_data_row']}",
        f"    colunas ({len(columns)}): "
        f"{json.dumps(columns, ensure_ascii=False)}",
    ]

    # Nomes gerados pelo extrator para cabeçalhos ausentes ou vazios.
    generated = any(
        name == f"coluna_{region['min_col'] + offset}"
        for offset, name in enumerate(columns)
    )

    return lines, generated


def _describe_xlsx(path: Path, aba: str | None) -> str:
    schema = get_schema(path)
    selected = _select_sheets(schema, aba)

    n_sheets = len(schema["sheets"])
    n_regions = sum(len(regions) for regions in selected.values())

    if len(selected) == n_sheets:
        lines = [
            f"{path.name} (XLSX): {_plural(n_sheets, 'aba', 'abas')}, "
            f"{_plural(n_regions, 'região tabular', 'regiões tabulares')}."
        ]
    else:
        todas = ", ".join(repr(title) for title in schema["sheets"])
        lines = [
            f"{path.name} (XLSX): exibindo só a aba {next(iter(selected))!r} "
            f"(o arquivo tem {n_sheets} abas: {todas})."
        ]

    any_generated = False

    for title, regions in selected.items():
        if not regions:
            lines.append(
                f"Aba {title!r}: nenhuma região tabular detectada "
                f"(aba vazia ou sem estrutura reconhecível)."
            )
            continue

        lines.append(
            f"Aba {title!r}: "
            f"{_plural(len(regions), 'região', 'regiões')}"
        )

        for index, region in enumerate(regions, start=1):
            region_lines, generated = _format_region(index, region)
            lines.extend(region_lines)
            any_generated = any_generated or generated

    if any_generated:
        lines.append(
            'Nota: "coluna_N" é nome gerado (cabeçalho ausente ou vazio); '
            "N é o número da coluna na aba."
        )

    lines.append(
        "Registros não lidos. Para valores das células, use create_json."
    )

    return "\n".join(lines)


def _describe_csv(path: Path, aba: str | None) -> str:
    size = path.stat().st_size
    size_text = f"{size} bytes" if size < 1024 else f"{size / 1024:.1f} KB"

    lines = [
        f"{path.name} (CSV, {size_text})",
        "CSV não tem abas nem regiões detectadas: não existe schema "
        "estrutural para este formato, e as colunas só podem ser "
        "conhecidas lendo o arquivo.",
    ]

    if aba:
        lines.append("O parâmetro aba foi ignorado: CSV não tem abas.")

    lines.append("Para ver colunas e valores, use create_json.")

    return "\n".join(lines)


@tool
def consultar_schema(name: str, ext: str, aba: str | None = None) -> str:
    """
    Consulta a ESTRUTURA de uma planilha (abas, regiões tabulares, linha
    do cabeçalho, limites e nomes das colunas) sem ler os registros.

    Use esta ferramenta ANTES de create_json sempre que a pergunta puder
    ser respondida só com metadados: quais abas existem, quais colunas uma
    aba tem, como a planilha está organizada, onde fica o cabeçalho, quais
    regiões foram detectadas, em quais abas existe uma coluna. Use
    create_json apenas quando precisar dos valores das células (filtrar,
    contar, somar ou listar registros).

    Para CSV não há abas nem regiões: a resposta informa isso.

    Args:
        name: Nome do arquivo sem extensão, por exemplo matriz.
        ext: Extensão do arquivo, csv ou xlsx.
        aba: Opcional. Nome de uma aba; se informado, retorna só ela.
    """
    ext = ext.lower().lstrip(".")

    if ext not in ("csv", "xlsx"):
        raise ValueError("Extensão não suportada")

    path = _resolve_path(name, ext)

    if ext == "csv":
        return _describe_csv(path, aba)

    return _describe_xlsx(path, aba)
